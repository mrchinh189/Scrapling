"""Xuất báo cáo DOCX Price Intelligence (bám mẫu v4) từ VIEW-MODEL dùng chung.

DOCX chỉ ĐỌC LẠI view-model, KHÔNG tự tính lại → bảo đảm DOCX = Web = Telegram.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app.config import DATA_DIR


def _add_hyperlink(paragraph, text: str, url: str | None):
    """Thêm text dạng hyperlink (nếu có url) vào paragraph."""
    from docx.oxml.shared import OxmlElement, qn

    if not url:
        paragraph.add_run(text + " (form nội bộ)")
        return
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    hl.append(run)
    paragraph._p.append(hl)


def build_intel_docx(vm: dict, report_id: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor

    meta = vm["meta"]
    doc = Document()

    # ① Trang bìa
    doc.add_heading("Báo cáo Price Intelligence — Giá NVL Masterbatch", level=0)
    doc.add_paragraph(f"Ngày giá: {meta['price_date']}  ·  "
                      f"Tỷ giá USD/VND: {meta['usd_vnd']:,.0f}  ·  "
                      f"Số NVL theo dõi: {meta['n_materials']}")
    doc.add_paragraph(f"Tạo lúc (giờ chạy pipeline): "
                      f"{datetime.now().strftime('%d/%m/%Y %H:%M')}")

    # ② Cách đọc
    doc.add_heading("① Cách đọc báo cáo", level=1)
    doc.add_paragraph(
        "4 câu hỏi mua hàng: (1) Giá đang ở đâu? (2) Vì sao biến động? "
        "(3) Nguồn nào tốt nhất khi quy về at-sight? (4) Nên mua hay chờ? "
        "Mẹo: cột 'At-sight tương đương' là con số xếp hạng công bằng giữa các "
        "điều khoản thanh toán khác nhau.")

    # ③ KPI tổng quan
    doc.add_heading("② Tổng quan nhanh (KPI)", level=1)
    t = doc.add_table(rows=1, cols=6)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(["NVL", "Giá (USD/tấn)", "%tuần", "At-sight (đ/kg)",
                           "Nguồn tốt nhất", "Ngày · Tươi"]):
        t.rows[0].cells[i].text = h
    for k in vm["kpis"]:
        c = t.add_row().cells
        c[0].text = k["label"]
        c[1].text = f"{k['latest_usd_ton']:,.0f}"
        pct = k.get("change_pct")
        c[2].text = f"{pct:+.2f}%" if pct is not None else "—"
        if pct is not None and c[2].paragraphs[0].runs:
            c[2].paragraphs[0].runs[0].font.color.rgb = (
                RGBColor(0xC0, 0, 0) if pct > 0 else RGBColor(0, 0x80, 0))
        c[3].text = f"{k['best_at_sight_vnd_kg']:,.0f}" if k.get("best_at_sight_vnd_kg") else "—"
        c[4].text = k.get("best_source", "—")
        c[5].text = f"{k['date']} {k['freshness']}"

    # ④ Phân tích tổng hợp (narrative)
    doc.add_heading("③ Phân tích tổng hợp", level=1)
    for line in vm["narrative"]["combined_md"].split("\n\n"):
        doc.add_paragraph(line.replace("**", ""))

    # ⑤ Quy đổi at-sight tương đương — từng NVL
    doc.add_heading("④ Quy đổi GIÁ CHUNG → At-sight tương đương", level=1)
    for prod, rows in vm["landed"].items():
        if not rows:
            continue
        doc.add_heading(prod, level=3)
        tb = doc.add_table(rows=1, cols=6)
        tb.style = "Light List Accent 1"
        for i, h in enumerate(["Khu vực", "Loại giá", "Thanh toán",
                               "Landed (đ/kg)", "At-sight (đ/kg)", "Nguồn"]):
            tb.rows[0].cells[i].text = h
        for r in sorted(rows, key=lambda x: (x["at_sight_equiv"] is None, x["at_sight_equiv"] or 0)):
            cc = tb.add_row().cells
            cc[0].text = r["region"]
            cc[1].text = r["price_type"]
            cc[2].text = r["payment_term"]
            cc[3].text = f"{r['landed_vnd_kg']:,.0f}" if r["landed_vnd_kg"] else "—"
            cc[4].text = (f"{r['at_sight_equiv']:,.0f}" + (" ⭐" if r["is_best"] else "")) \
                if r["at_sight_equiv"] else "—"
            cc[5].text = r["source"]

    # ⑦ Cảnh báo & đề xuất
    doc.add_heading("⑤ Cảnh báo & Đề xuất hành động", level=1)
    for a in vm["alerts"]:
        p = doc.add_paragraph()
        p.add_run(f"[{a['severity']}] {a['title']}\n").bold = True
        p.add_run(f"{a['detail']}\n{a['recommendation']}")

    # ⑨ Dự báo
    doc.add_heading("⑥ Dự báo (baseline)", level=1)
    for prod, fc in vm["forecasts"].items():
        doc.add_paragraph(
            f"{prod}: model {fc['model']} · tin cậy {fc['confidence']} · "
            f"Theil's U {fc['theils_u']} · cơ sở: {fc['basis']}")
        end = fc["points"][-1]
        doc.add_paragraph(
            f"   T+{fc['horizon']}: {end['yhat']:,.0f} "
            f"(khoảng {end['lower']:,.0f}–{end['upper']:,.0f})", style="List Bullet")

    # Nguồn & độ tươi (kèm link + ngày giá)
    doc.add_heading("⑦ Nguồn & độ tươi", level=1)
    for s in vm["sources"]:
        p = doc.add_paragraph(style="List Bullet")
        _add_hyperlink(p, s["label"], s["url"])
        p.add_run(f" — ngày giá {s['latest_date']} {s['freshness']}")

    out_dir = DATA_DIR / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{report_id}.docx"
    doc.save(path)
    return str(path)
