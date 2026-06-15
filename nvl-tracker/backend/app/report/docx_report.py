"""Xuất báo cáo phân tích ra file .docx."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app.config import DATA_DIR
from app.models import PriceChange


def build_docx(
    title: str,
    analysis_md: str,
    changes: list[PriceChange],
    report_id: str,
) -> str:
    """Tạo file .docx từ nội dung phân tích + bảng giá. Trả về đường dẫn file."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Ngày tạo: {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    # Bảng giá
    doc.add_heading("Bảng giá NVL", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["NVL", "Đơn vị", "Giá hiện tại", "Giá trước", "Thay đổi %"]):
        hdr[i].text = h
    for c in changes:
        row = table.add_row().cells
        row[0].text = f"{c.material_name} ({c.material_code})"
        row[1].text = c.unit
        row[2].text = f"{c.current_price:,.0f}"
        row[3].text = f"{c.previous_price:,.0f}" if c.previous_price is not None else "—"
        pct_cell = row[4]
        pct_cell.text = f"{c.change_pct:+.2f}%" if c.change_pct is not None else "—"
        if c.change_pct is not None:
            run = pct_cell.paragraphs[0].runs[0]
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00) if c.change_pct > 0 else RGBColor(0x00, 0x80, 0x00)

    # Nội dung phân tích (Markdown đơn giản -> Word)
    doc.add_heading("Phân tích", level=1)
    _render_markdown(doc, analysis_md, Pt)

    out_dir = DATA_DIR / "reports"
    out_dir.mkdir(exist_ok=True)
    path = out_dir / f"{report_id}.docx"
    doc.save(path)
    return str(path)


def _render_markdown(doc, md: str, Pt) -> None:
    """Chuyển Markdown cơ bản (heading, gạch đầu dòng, bảng bỏ qua) sang đoạn Word."""
    for line in md.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("|"):
            continue  # bảng đã render riêng ở trên
        elif stripped.startswith(("- ", "* ")):
            doc.add_paragraph(stripped[2:].replace("**", ""), style="List Bullet")
        else:
            doc.add_paragraph(stripped.replace("**", ""))


__all__ = ["build_docx"]
